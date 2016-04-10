# -*- coding: utf-8 -*-
"""
Created on Tue Mar  8 17:49:47 2016

@author: dell
"""

import pandas as pd
from openpyxl import load_workbook
import random
import itertools
import json
   
recipeID=0
def read_moulds (category):
    
    filename = 'App Data/Recipe Moulds/mould_'+category+'.txt'
    #mould_file = open(filename, 'r', encoding = "utf8")
    mould_file = open(filename, 'r')
   
    #print (mould_file.readline())
    no_mould_sections  = int (mould_file.readline())
    
    if no_mould_sections ==1:
         # read first line containing placeholders
        list_placeholders = mould_file.readline().strip().split(',') #placeholder = relabeled category according to the mould
    
        
        #this contains the max allowed limit for each category of ingredient      
        list_max_limits_of_placeholders = mould_file.readline().strip().split(',')
        recipe_procedure_lines =   mould_file.readlines()
        return (list_placeholders, list_max_limits_of_placeholders, recipe_procedure_lines)
    
    else:
        #print ('i am in else')
        section = []
        l=[]
        for line in mould_file.readlines():
            #print (line)
            if line.strip('\n'):
                section.append(line)
            else:
                section= ''.join(section)
                l.append(section)
                section = []
        # yield any remaining lines as a section too
        if section:
            section= ''.join(section)
            l.append(section)
            
        
        r = random.randint(0, no_mould_sections-1)
        random_mould_section_selected = l[r]
        
        lines = random_mould_section_selected.split ('\n')
        list_placeholders = lines[0].strip().split(',')
        list_max_limits_of_placeholders = lines[1].strip().split(',')
        recipe_procedure_lines = lines[2 : len(lines)-1]
        return (list_placeholders, list_max_limits_of_placeholders, recipe_procedure_lines)
    
def app_get_and_print_the_recipe(category, r_rowNum):
    wb = pd.ExcelFile("App Data/Database of Stored Recipes.xlsx")
    WS_RecipeID_IngreIDs= wb.parse(category)
    
    wb1 = pd.ExcelFile("App Data/Relabeled_Ingredient_List.xlsx")
    WS_IngreID_IngreName_Category= wb1.parse(category)
    
    (list_placeholders, list_max_limits_of_placeholders, recipe_procedure_lines) = read_moulds (category)    
    dict_cat_ingreNames = dict()
    
    #read a random recipe excel row from the 'Database of Stored Recipes'
    if r_rowNum == -1 :
        r_rowNum = random.randint(0, len (WS_RecipeID_IngreIDs.index)-1) # we want [a,b)...  so do [a, b-1] 
    
    selected_recipeID = WS_RecipeID_IngreIDs.iloc[r_rowNum]['Recipe ID']
    string_of_ingre_ids = WS_RecipeID_IngreIDs.iloc[r_rowNum]['Ingredient ID']
    
    list_selected_recipeIngreIDs = string_of_ingre_ids.split(',')
    for ingreID in list_selected_recipeIngreIDs:
        ingreID = ingreID.replace('[', '') #to extract only the ingre ID and ignore the rest
        ingreID = ingreID.replace(']', '')
        ingreID = int (ingreID) #convert string to int 
        
        for z in range (0, len(WS_IngreID_IngreName_Category.index)):
            if ingreID == WS_IngreID_IngreName_Category.iloc[z]['Ingredient ID']:
                ingreName = WS_IngreID_IngreName_Category.iloc[z]['Ingredient Name']
                placeholder = WS_IngreID_IngreName_Category.iloc[z]['Category']
                #print (placeholder)
                break
            
        if placeholder not in dict_cat_ingreNames:
            dict_cat_ingreNames[placeholder] = [ingreName]
        else:
            dict_cat_ingreNames[placeholder].append(ingreName)
    
    index=0
    print_recipe=[]
    print_ingred_list=[]
    for line in recipe_procedure_lines: 
            num_of_stars_found_in_a_line= line.count('*')
            
            if num_of_stars_found_in_a_line >0:
                
                for i in range (0, num_of_stars_found_in_a_line):
                    index_of_star_to_be_replaced = line.find('*')
                    
                    placeholder = list_placeholders[index] #placeholder = fruit, vegetable, oil, etc.

                    line = list(line)
                    tempstring = ''
                    for el in dict_cat_ingreNames[placeholder]:
                        tempstring = (str(el)) + ', '
                    tempstring = tempstring[:-2]
                    line[index_of_star_to_be_replaced] = str(tempstring)
                    print_ingred_list.append(dict_cat_ingreNames[placeholder])
                    line = "".join(line)                                      
                    index = index+1 #increment the placeholder index to replace the next * appropriately
                    
                print_recipe.append(line) 
           
            else:
                print_recipe.append(line)

    
    combined = list(itertools.chain.from_iterable(print_ingred_list))

    cluster_name = "\"cluster\" : \"" + category + "\""
    r_rowNum = category + str(r_rowNum)
    recipe_id = "\"recipeID\" : " + str(r_rowNum) + ""
    ingredients = []
    recipe = []
    for ingre in combined : 
        ingredients.append({"ingredient" : ingre})
    for step in print_recipe :
        recipe.append({"step" : step})

    ingredients =  "\"ingredients\" :" + json.dumps(ingredients)
    recipe =  "\"recipe\" :" +json.dumps(recipe) 

    #print (cluster_name)
    #print (recipe_id)
    #print (ingredients)
    #print (recipe)
    return ("{"+cluster_name + ", " + recipe_id + ", " + ingredients + ", " + recipe + "}")

def get_all_clusters():
    clusters = []
    clusters.append({"cluster" : "idli"})
    clusters.append({"cluster" : "roti_naan"})
    clusters.append({"cluster" : "pakode_bhajia"})
    clusters.append({"cluster" : "khichdi"})
    return "{" + "\"clusters\" : " + json.dumps(clusters) + "}"

#app_get_and_print_the_recipe('idli',-1)